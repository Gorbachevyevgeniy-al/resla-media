# Публикует один слот из autopost/slots.json в Instagram @cosmetics.kaz.kz. Не зависит от /mnt/project-files:
# расписание, подписи и медиа берутся из этого публичного репозитория. Маркер берётся из переменной IG_ACCESS_TOKEN.
# Запуск в любой облачной сессии:
#   curl -fsSL https://raw.githubusercontent.com/Gorbachevyevgeniy-al/resla-media/main/autopost/run_slot.py -o /tmp/run_slot.py && python3 /tmp/run_slot.py <id> [<id> ...]
# kind: post | story | reel | carousel (у карусели url — список картинок). Несколько id публикуются по очереди с паузой 60 с (пачки сториз). Уже вышедший пост/рилс (та же подпись) не публикуется второй раз.
import datetime, json, os, sys, time, urllib.parse, urllib.request
RAW='https://raw.githubusercontent.com/Gorbachevyevgeniy-al/resla-media/main/autopost/slots.json'
B='https://graph.instagram.com/v23.0'; U='17841461175377751'; T=os.environ['IG_ACCESS_TOKEN']
XLSX='/mnt/project-files/catalog/Sinep_прайс-лист_Kaspi.xlsx'

class Fail(Exception): pass

def call(method, path, **p):
    p['access_token']=T; data=urllib.parse.urlencode(p).encode()
    url=f'{B}/{path}' + ('' if method=='POST' else '?'+data.decode())
    req=urllib.request.Request(url, data=data if method=='POST' else None, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as r: return json.load(r)
    except urllib.error.HTTPError as e: raise Fail(f'HTTP {e.code}: {e.read().decode()}')

def already_out(cap):
    head=cap.strip()[:120]
    for m in call('GET', 'me/media', fields='caption,permalink,timestamp', limit='15').get('data', []):
        if (m.get('caption') or '').strip()[:120]==head: return m
    return None

def wait(c, n):
    for _ in range(n):
        st=call('GET', c, fields='status_code')['status_code']
        if st=='FINISHED': return
        if st in ('ERROR','EXPIRED'): raise Fail(f'container {c}: {st}')
        time.sleep(5)

def publish(s):
    cap=s.get('caption_text')
    if cap and s['kind']!='story':
        m=already_out(cap)
        if m: return {'skipped':'already published', **m}
    if s['kind']=='reel':
        p={'media_type':'REELS','video_url':s['url'],'share_to_feed':'true','thumb_offset':str(s['thumb'] if s.get('thumb') is not None else 1000)}
    elif s['kind']=='carousel':
        # Карусель: url — список картинок (2–10), каждая отдельным контейнером is_carousel_item, затем общий контейнер CAROUSEL
        kids=[call('POST', f'{U}/media', image_url=u, is_carousel_item='true')['id'] for u in s['url']]
        for k in kids: wait(k, 30)
        p={'media_type':'CAROUSEL','children':','.join(kids)}
    else:
        p={'image_url':s['url']}
        if s['kind']=='story': p['media_type']='STORIES'
    if cap: p['caption']=cap
    c=call('POST', f'{U}/media', **p)['id']
    wait(c, 120 if s['kind']=='reel' else 30)
    # Instagram иногда отвечает «Media ID is not available» (код 9007) сразу после FINISHED: ждём и пробуем ещё
    for i in range(6):
        try: m=call('POST', f'{U}/media_publish', creation_id=c)['id']; break
        except Fail as e:
            if '9007' not in str(e) or i==5: raise
            time.sleep(10)
    return call('GET', m, fields='id,permalink,timestamp,media_product_type')

def journal(s, res):
    # Лист «Журнал постов» есть только там, где подключена общая папка проекта; без неё просто пропускаем.
    if not os.path.exists(XLSX) or res.get('skipped'): return 'journal skipped'
    import openpyxl
    wb=openpyxl.load_workbook(XLSX); ws=wb['Журнал постов']
    row=next(i for i in range(2,ws.max_row+2) if ws.cell(i,1).value is None)
    d=datetime.datetime.fromisoformat(s['utc'].replace('Z','+00:00')).astimezone(datetime.timezone(datetime.timedelta(hours=5)))
    for c,v in [(1,datetime.datetime(d.year,d.month,d.day)),(2,s['fmt']),(3,s['code']),(5,s['theme']),(6,res.get('permalink')),(7,'Опубликован')]:
        ws.cell(row,c).value=v
    wb.save(XLSX); return f'journal row {row}'

local=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'slots.json')
slots=json.load(open(local, encoding='utf-8')) if os.path.exists(local) else json.load(urllib.request.urlopen(RAW+f'?t={int(time.time())}', timeout=60))
fail=[]
for i,sid in enumerate(sys.argv[1:]):
    if i: time.sleep(60)
    try:
        res=publish(slots[sid]); print(sid, 'OK', json.dumps(res, ensure_ascii=False), journal(slots[sid], res), flush=True)
    except Exception as e:
        print(sid, 'FAILED', str(e)[-300:], flush=True); fail.append(sid)
print('DONE', len(sys.argv)-1-len(fail), 'of', len(sys.argv)-1, 'failed:', fail)
sys.exit(1 if fail else 0)
